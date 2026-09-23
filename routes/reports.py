from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from datetime import datetime, timedelta
from database import one, query, fmt_date
import pandas as pd, io, calendar, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.units import cm

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

def _range(period):
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    if period == 'today': return today, today + timedelta(days=1)
    if period == 'week':
        s = today - timedelta(days=today.weekday()); return s, s + timedelta(days=7)
    if period == 'month':
        s = today.replace(day=1); _, ld = calendar.monthrange(today.year, today.month)
        return s, today.replace(day=ld, hour=23, minute=59, second=59)
    return None, None

@reports_bp.route('/')
@login_required
def index():
    return render_template('reports/index.html')

@reports_bp.route('/collection')
@login_required
def collection_report():
    period = request.args.get('period', 'today')
    start_str = request.args.get('start_date', '').strip()
    end_str   = request.args.get('end_date',   '').strip()
    start = end = None

    if period == 'custom' and start_str and end_str:
        try:
            start = datetime.strptime(start_str, '%Y-%m-%d')
            end   = datetime.strptime(end_str, '%Y-%m-%d').replace(hour=23, minute=59, second=59)
        except ValueError:
            flash('Invalid date format. Please use the date picker.', 'warning')
            start = end = None

    if start is None or end is None:
        start, end = _range(period if period != 'custom' else 'today')
        if not end:
            start = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            end   = start + timedelta(days=1)

    rows = query("SELECT * FROM payments WHERE DATE(collection_date)>=? AND DATE(collection_date)<=? ORDER BY collection_date DESC",
                 (start.strftime('%Y-%m-%d'), end.strftime('%Y-%m-%d')))
    for p in rows: p['collection_date_str'] = fmt_date(p.get('collection_date'))
    total_collected = sum(p['paid_amount'] for p in rows)
    total_pending   = sum(p['pending_amount'] for p in rows)
    return render_template('reports/collection.html',
                           payments=rows, period=period, start=start, end=end,
                           total_collected=round(total_collected,2),
                           total_pending=round(total_pending,2),
                           start_str=start_str, end_str=end_str)

@reports_bp.route('/loan-report')
@login_required
def loan_report():
    status = request.args.get('status', '')
    loans  = query("SELECT * FROM loans WHERE loan_status=? ORDER BY id" if status else "SELECT * FROM loans ORDER BY id",
                   (status,) if status else ())
    for l in loans:
        l['loan_date'] = fmt_date(l.get('loan_date'))
    total_issued = sum(l['loan_amount'] for l in loans)
    total_paid   = sum(l['paid_amount'] for l in loans)
    return render_template('reports/loan_report.html', loans=loans, status=status,
                           total_issued=round(total_issued,2), total_collected=round(total_paid,2))

@reports_bp.route('/pending-report')
@login_required
def pending_report():
    rows = query("""
        SELECT p.loan_id, p.customer_name, p.pending_amount AS pending,
               p.collection_date AS last_date, l.phone
        FROM payments p JOIN loans l ON l.loan_id=p.loan_id
        WHERE p.id IN (
            SELECT id FROM payments p2 WHERE p2.loan_id=p.loan_id
            ORDER BY p2.collection_date DESC LIMIT 1
        ) AND p.pending_amount > 0 ORDER BY p.pending_amount DESC
    """)
    for r in rows: r['last_date_str'] = fmt_date(r.get('last_date'))
    return render_template('reports/pending_report.html', rows=rows,
                           total_pending=round(sum(r['pending'] for r in rows),2))

# ── Excel exports ─────────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill("solid", start_color="0F4C81")
HEADER_FONT   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
DATA_FONT     = Font(name="Arial", size=9)
ALT_FILL      = PatternFill("solid", start_color="EEF3FA")
THIN          = Side(style='thin', color="CCCCCC")
BORDER        = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER        = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT          = Alignment(horizontal='left',  vertical='center', wrap_text=True)

def _style_sheet(ws, headers, col_widths):
    for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill = HEADER_FILL; cell.font = HEADER_FONT
        cell.border = BORDER;    cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[1].height = 30

def _apply_data_rows(ws, rows_data, start_row=2):
    for ri, row in enumerate(rows_data, start_row):
        fill = ALT_FILL if ri % 2 == 0 else PatternFill()
        for ci, val in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.font = DATA_FONT; c.border = BORDER
            c.fill = fill
            c.alignment = LEFT

@reports_bp.route('/export/excel/loans')
@login_required
def export_excel_loans():
    import openpyxl
    loans = query("SELECT * FROM loans ORDER BY id")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Loans"
    ws.freeze_panes = "A2"
    headers = ["Sl.No","Loan No","Loan Date","Name","Mobile Number","Ref","Due Date",
               "Follower","Loan Amount","Document Charge","Principle Amount","Interest Amount",
               "Agreement Value","Tenure","Due Amount","Paid Due","Paid Amount",
               "Feature Dues","Balance Agreement","Int %","Rate %/mo",
               "Vehicle Model","Vehicle Number","Document Status","Remarks",
               "Penalty Amount","Penalty Description","Status"]
    widths  = [6,10,12,22,22,10,10,12,13,15,15,15,15,8,12,9,13,12,18,8,9,24,16,18,24,13,24,10]
    _style_sheet(ws, headers, widths)
    rows_data = []
    for i, l in enumerate(loans, 1):
        rows_data.append([
            i, l['loan_id'], fmt_date(l.get('loan_date'), '%d-%m-%Y'),
            l['customer_name'], l['phone'], l['ref'], l['due_day'],
            l['follower'], l['loan_amount'], l['document_charge'], l['principle_amount'],
            l['interest_amount'], l['agreement_value'], l['tenure'],
            l['due_amount'], l['paid_due'], l['paid_amount'],
            l['feature_dues'], l['balance_agreement'], l['int_pct'],
            l.get('interest_rate', 0),
            l['vehicle_model'], l['vehicle_number'], l['document_status'],
            l['remarks'], l['penalty_amount'], l['penalty_description'], l['loan_status'],
        ])
    _apply_data_rows(ws, rows_data)
    # Totals row
    tr = len(loans) + 2
    ws.cell(tr, 1, "TOTAL").font = Font(bold=True, name="Arial")
    for col, field in [(9,'loan_amount'),(10,'document_charge'),(11,'principle_amount'),
                       (12,'interest_amount'),(13,'agreement_value'),(15,'due_amount'),
                       (17,'paid_amount'),(19,'balance_agreement'),(26,'penalty_amount')]:
        ws.cell(tr, col, sum(l[field] for l in loans)).font = Font(bold=True, name="Arial", size=9)
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='agaram_loans.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@reports_bp.route('/export/excel/customers')
@login_required
def export_excel_customers():
    import openpyxl
    from routes.customers import _mask_aadhaar
    data = query("SELECT * FROM customers WHERE is_active=1 ORDER BY name")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Customers"; ws.freeze_panes = "A2"
    headers = ["Name","Phone","Alt Phone","Address","Aadhaar","Occupation","Date Added"]
    widths  = [24,16,16,30,16,16,14]
    _style_sheet(ws, headers, widths)
    # Aadhaar is masked for non-admins in bulk exports too (audit item 31)
    aadhaar_of = (lambda d: d['aadhaar']) if current_user.is_admin else (lambda d: _mask_aadhaar(d['aadhaar']))
    _apply_data_rows(ws, [[d['name'],d['phone'],d['alt_phone'],d['address'],
                           aadhaar_of(d),d['occupation'],fmt_date(d['date_added'],'%d-%m-%Y')] for d in data])
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='agaram_customers.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@reports_bp.route('/export/excel/payments')
@login_required
def export_excel_payments():
    import openpyxl
    start = request.args.get('start', '')
    end   = request.args.get('end', '')
    if start and end:
        data = query("SELECT * FROM payments WHERE collection_date>=? AND collection_date<=? ORDER BY collection_date DESC",
                     (start, end + ' 23:59:59'))
    else:
        data = query("SELECT * FROM payments ORDER BY collection_date DESC")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Payments"; ws.freeze_panes = "A2"
    headers = ["Payment ID","Loan ID","Customer","Date","EMI","Carry Fwd","Total Due",
               "Paid","Pending","Penalty","Month#","Remarks","By"]
    widths  = [14,10,22,12,12,12,12,12,12,12,7,24,12]
    _style_sheet(ws, headers, widths)
    _apply_data_rows(ws, [[d['payment_id'],d['loan_id'],d['customer_name'],
                           fmt_date(d['collection_date'],'%d-%m-%Y'),
                           d['emi_amount'],d['carry_forward'],d['total_due'],
                           d['paid_amount'],d['pending_amount'],d['penalty_amount'],
                           d['month_number'],d['collection_remarks'],d['collected_by']] for d in data])
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='agaram_payments.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@reports_bp.route('/export/excel/pending')
@login_required
def export_excel_pending():
    import openpyxl
    rows = query("""
        SELECT p.loan_id, p.customer_name, p.pending_amount, p.collection_date, l.phone, l.due_amount
        FROM payments p JOIN loans l ON l.loan_id=p.loan_id
        WHERE p.id IN (SELECT id FROM payments p2 WHERE p2.loan_id=p.loan_id ORDER BY p2.collection_date DESC LIMIT 1)
        AND p.pending_amount > 0 ORDER BY p.pending_amount DESC""")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Pending"; ws.freeze_panes = "A2"
    _style_sheet(ws, ["Loan ID","Customer","Phone","Monthly Due","Pending","Last Payment"], [10,24,16,13,13,14])
    _apply_data_rows(ws, [[r['loan_id'],r['customer_name'],r['phone'],r['due_amount'],
                           r['pending_amount'],fmt_date(r['collection_date'],'%d-%m-%Y')] for r in rows])
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name='agaram_pending.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ── PDF export ────────────────────────────────────────────────────────────────
@reports_bp.route('/export/pdf/<report_type>')
@login_required
def export_pdf(report_type):
    output = io.BytesIO()
    doc    = SimpleDocTemplate(output, pagesize=landscape(A4),
                               topMargin=1*cm,bottomMargin=1*cm,leftMargin=1.5*cm,rightMargin=1.5*cm)
    styles = getSampleStyleSheet()
    ts = ParagraphStyle('T', parent=styles['Heading1'], fontSize=14,
                        textColor=colors.HexColor('#0F4C81'), spaceAfter=10)
    tbl_style = TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#0F4C81')),
        ('TEXTCOLOR', (0,0),(-1,0),colors.white),
        ('FONTNAME',  (0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',  (0,0),(-1,0),8),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#EEF3FA')]),
        ('FONTSIZE',  (0,1),(-1,-1),7),
        ('GRID',      (0,0),(-1,-1),0.4,colors.HexColor('#CCCCCC')),
        ('PADDING',   (0,0),(-1,-1),4),
    ])
    elements = []
    if report_type == 'loans':
        elements += [Paragraph('Agaram Finance – Loan Report', ts),
                     Paragraph(f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")}', styles['Normal']),
                     Spacer(1,0.3*cm)]
        data  = query("SELECT * FROM loans ORDER BY id")
        tbl   = [['#','Loan No','Name','Loan Amt','Int%','Due Amt','Paid','Balance','Vehicle No','Doc Status','Status']]
        for i,d in enumerate(data,1):
            tbl.append([i,d['loan_id'],(d['customer_name'] or '')[:18],
                        f"₹{d['loan_amount']:,.0f}",d['int_pct'],f"₹{d['due_amount']:,.0f}",
                        f"₹{d['paid_amount']:,.0f}",f"₹{d['balance_agreement']:,.0f}",
                        d['vehicle_number'],(d['document_status'] or '')[:14],d['loan_status']])
        fname = 'loans_report.pdf'
    elif report_type == 'collection':
        elements += [Paragraph('Agaram Finance – Collection Report', ts),
                     Paragraph(f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")}', styles['Normal']),
                     Spacer(1,0.3*cm)]
        data  = query("SELECT * FROM payments ORDER BY collection_date DESC LIMIT 200")
        tbl   = [['Payment ID','Loan ID','Customer','Date','EMI','Carry Fwd','Paid','Pending','Penalty']]
        for d in data:
            tbl.append([d['payment_id'],d['loan_id'],(d['customer_name'] or '')[:18],
                        fmt_date(d['collection_date'],'%d-%m-%Y'),
                        f"₹{d['emi_amount']:,.0f}",f"₹{d['carry_forward']:,.0f}",
                        f"₹{d['paid_amount']:,.0f}",f"₹{d['pending_amount']:,.0f}",
                        f"₹{d['penalty_amount']:,.0f}"])
        fname = 'collection_report.pdf'
    else:
        return redirect(url_for('reports.index'))
    table = Table(tbl); table.setStyle(tbl_style)
    elements.append(table); doc.build(elements); output.seek(0)
    return send_file(output, as_attachment=True, download_name=fname, mimetype='application/pdf')
